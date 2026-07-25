"""Shared openpyxl workbook helpers for MIS exports."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Iterable, Sequence

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from setup.models import OrganizationSettings

HEADER_FILL = PatternFill("solid", fgColor="0F1C2E")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=14, color="0F1C2E")
META_FONT = Font(name="Calibri", size=10, color="5B6B7C")
CELL_FONT = Font(name="Calibri", size=10)
MONEY_FONT = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
ALT_FILL = PatternFill("solid", fgColor="F4F6F9")


def _cell_value(value: Any):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, date):
        return value
    return value


def new_workbook(*, title: str, subtitle: str = "") -> tuple[Workbook, Worksheet, int]:
    """Create a workbook with company header. Returns (wb, ws, next_row)."""
    org = OrganizationSettings.load()
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = org.display_name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, name="Calibri", size=12)
    meta_bits = [
        subtitle.strip(),
        f"Country: {org.country}" if org.country else "",
        f"Generated: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}",
    ]
    ws["A3"] = " · ".join(b for b in meta_bits if b)
    ws["A3"].font = META_FONT
    return wb, ws, 5


def write_table(
    ws: Worksheet,
    *,
    start_row: int,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    money_cols: Sequence[int] | None = None,
    freeze: bool = True,
) -> int:
    """Write a header + data table. Returns the next empty row index."""
    money_cols = set(money_cols or ())
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN

    row_idx = start_row
    for i, row in enumerate(rows):
        row_idx = start_row + 1 + i
        for col, raw in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=_cell_value(raw))
            cell.font = CELL_FONT
            cell.border = THIN
            if i % 2 == 1:
                cell.fill = ALT_FILL
            if col in money_cols and isinstance(raw, (Decimal, int, float)) and raw != "":
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(raw, date) and not isinstance(raw, datetime):
                cell.number_format = "YYYY-MM-DD"

    last_row = max(row_idx, start_row)
    last_col = len(headers)
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(last_col)}{last_row}"
    _autosize(ws, last_col, last_row)
    return last_row + 2


def add_sheet(wb: Workbook, title: str) -> Worksheet:
    name = title[:31]
    if name in wb.sheetnames:
        n = 2
        while f"{name[:28]}_{n}" in wb.sheetnames:
            n += 1
        name = f"{name[:28]}_{n}"
    return wb.create_sheet(name)


def _autosize(ws: Worksheet, last_col: int, last_row: int) -> None:
    for col in range(1, last_col + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(last_row, 200) + 1):
            val = ws.cell(row=row, column=col).value
            if val is None:
                continue
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def workbook_response(wb: Workbook, filename: str) -> HttpResponse:
    if not filename.endswith(".xlsx"):
        filename = f"{filename}.xlsx"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
